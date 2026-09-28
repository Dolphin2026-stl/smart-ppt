import sys,tempfile,unittest,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from template_engine import generate
from common import geometry
class FullDeckTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);self.source=self.root/'source.pptx';self.out=self.root/'out.pptx'
  p=Presentation()
  for i in range(4):
   s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text=f'Original {i}';s.placeholders[1].text=f'Body {i}';s.notes_slide.notes_text_frame.text=f'Notes {i}'
  data=CategoryChartData();data.categories=['A','B'];data.add_series('Values',[1,2]);p.slides[2].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(2),Inches(4),Inches(3),data)
  p.save(self.source)
 def test_default_is_exact_complete_copy(self):
  generate(self.source,{},self.out);self.assertEqual(self.source.read_bytes(),self.out.read_bytes());self.assertEqual(len(Presentation(self.out).slides),4)
 def test_edit_preserves_other_parts_and_all_pages(self):
  p=Presentation(self.source);sid=str(p.slides[0].shapes.title.shape_id);box=geometry(p.slides[0].shapes.title)
  generate(self.source,{'edits':[{'page':0,'texts':{sid:'Changed'}}]},self.out)
  p=Presentation(self.out);self.assertEqual(len(p.slides),4);self.assertEqual(p.slides[0].shapes.title.text,'Changed');self.assertEqual(geometry(p.slides[0].shapes.title),box)
  with zipfile.ZipFile(self.source) as a,zipfile.ZipFile(self.out) as b:
   self.assertEqual(set(a.namelist()),set(b.namelist()))
   self.assertEqual([k for k in a.namelist() if a.read(k)!=b.read(k)],['ppt/slides/slide1.xml'])
 def test_explicit_delete_only(self):
  generate(self.source,{'delete_pages':[1]},self.out);p=Presentation(self.out);self.assertEqual([s.shapes.title.text for s in p.slides],['Original 0','Original 2','Original 3'])
  with zipfile.ZipFile(self.out) as z:self.assertNotIn('ppt/slides/slide2.xml',z.namelist())
 def test_duplicate_chart_notes_and_independent_text(self):
  sid=str(Presentation(self.source).slides[2].shapes.title.shape_id)
  generate(self.source,{'sequence':[0,1,{'source_page':2,'texts':{sid:'First'}},2,3]},self.out)
  p=Presentation(self.out);self.assertEqual(len(p.slides),5);self.assertEqual(p.slides[2].shapes.title.text,'First');self.assertEqual(p.slides[3].shapes.title.text,'Original 2')
  a=next(s.chart for s in p.slides[2].shapes if s.has_chart);b=next(s.chart for s in p.slides[3].shapes if s.has_chart)
  self.assertNotEqual(str(a.part.partname),str(b.part.partname));self.assertEqual(list(a.series[0].values),list(b.series[0].values));self.assertIn('Notes 2',p.slides[3].notes_slide.notes_text_frame.text)
 def test_legacy_extraction_requires_explicit_mode(self):
  with self.assertRaises(ValueError):generate(self.source,{'slides':[{'title':'wrong'}]},self.out)
if __name__=='__main__':unittest.main()
