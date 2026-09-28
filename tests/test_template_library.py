import sys,unittest,tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from pptx import Presentation
import template_library as lib
class LibraryTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);previous=lib.ROOT;lib.ROOT=self.root;self.addCleanup(setattr,lib,'ROOT',previous)
  for folder in ['school','company/nested']:
   p=self.root/'templates'/folder;p.mkdir(parents=True);deck=Presentation()
   for i in range(4):deck.slides.add_slide(deck.slide_layouts[1])
   deck.save(p/'complete.pptx')
 def test_all_nested_files_are_complete_candidates(self):
  items=lib.index_library();self.assertEqual(len(items),2);self.assertEqual([x['slides'] for x in items],[4,4]);self.assertEqual(len({x['id'] for x in items}),2)
 def test_select_keeps_every_byte(self):
  item=lib.index_library()[0];out=self.root/'work.pptx';lib.select(item['id'],out);self.assertEqual(out.read_bytes(),(self.root/item['path']).read_bytes())
if __name__=='__main__':unittest.main()
