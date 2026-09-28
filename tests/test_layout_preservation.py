"""Regression tests for actual layout invariants, pagination and negative cases."""
import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
from common import *
from style_engine import build_template,load_style,create_deck
from template_engine import generate
from validate_output import validate
from template_creator import create
from layout_inspector import inspect

class PreservationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        self.template=self.root/'template.pptx';build_template(load_style('minimal'),self.template)
        self.pic=self.root/'wide.png';Image.new('RGB',(900,300),'navy').save(self.pic)
    def run_plan(self,plan):
        plan.setdefault('mode','layouts');out=self.root/'out.pptx';generate(self.template,plan,out);return Presentation(out),out
    def assert_geometry(self,prs):
        for slide in prs.slides:
            for ph in slide.placeholders:
                target=slide.slide_layout.placeholders.get(ph.placeholder_format.idx)
                for a,b in zip(geometry(ph),geometry(target)):self.assertLessEqual(abs(a-b),1)
            self.assertFalse(any(s.shape_type==MSO_SHAPE_TYPE.TEXT_BOX for s in slide.shapes))
    def test_multiple_layouts_preserve_geometry(self):
        prs,out=self.run_plan({'allowed_layouts':[[0,0],[0,3]],'slides':[{'layout':0,'title':'Title','body':['subtitle']},{'layout':3,'title':'Two','body':['Left','Right']}]})
        self.assert_geometry(prs);self.assertTrue(validate(out)['passed']);self.assertEqual(len(prs.slides),2)
    def test_picture_frame_and_crop_preserved(self):
        prs,out=self.run_plan({'slides':[{'layout':5,'title':'Image','images':[str(self.pic)]}]})
        self.assert_geometry(prs);pic=next(s for s in prs.slides[0].shapes if hasattr(s,'image'))
        ph=prs.slides[0].slide_layout.placeholders.get(pic.placeholder_format.idx)
        self.assertAlmostEqual(pic.width/pic.height,ph.width/ph.height)
        self.assertFalse(any(x['code']=='picture_aspect' for x in validate(out)['issues']))
    def test_excess_slots_paginate(self):
        prs,_=self.run_plan({'slides':[{'layout':2,'title':'Paged','body':['one','two','three']}]})
        self.assertEqual(len(prs.slides),3);self.assert_geometry(prs)
    def test_long_text_preserves_all_characters(self):
        text='研究验证'*100
        prs,_=self.run_plan({'slides':[{'layout':2,'title':'Paged','body':[text],'max_chars':45}]})
        recovered=''.join(s.placeholders[1].text for s in prs.slides)
        self.assertEqual(text,recovered);self.assertGreater(len(prs.slides),1);self.assert_geometry(prs)
    def test_whitelist_rejects_other_layout(self):
        with self.assertRaises(ValueError):self.run_plan({'allowed_layouts':[[0,0]],'slides':[{'layout':2,'title':'No'}]})
    def test_auto_selection_only_allowed(self):
        prs,_=self.run_plan({'allowed_layouts':[[0,3]],'slides':[{'title':'Auto','body':['a','b']}]})
        self.assertEqual(prs.slides[0].slide_layout.name,'two-column')
    def test_insufficient_type_rejected(self):
        with self.assertRaises(ValueError):self.run_plan({'slides':[{'layout':0,'images':[str(self.pic)]}]})
    def test_source_unchanged(self):
        before=self.template.read_bytes();self.run_plan({'slides':[{'layout':0,'title':'OK'}]});self.assertEqual(before,self.template.read_bytes())
        with self.assertRaises(ValueError):generate(self.template,{'slides':[{}]},self.template)
    def test_negative_shift_detected(self):
        prs,out=self.run_plan({'slides':[{'layout':0,'title':'x'}]});prs.slides[0].placeholders[0].left+=91440;prs.save(out)
        self.assertFalse(validate(out)['passed'])
    def test_sample_page_clone_preserves_shapes(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);box=s.shapes.add_textbox(914400,914400,3657600,1828800);box.text='old';s.shapes.add_picture(str(self.pic),914400,3657600,width=3657600)
        p.save(self.template)
        before={str(x.shape_id):geometry(x) for x in s.shapes}
        prs,out=self.run_plan({'mode':'sample-pages','allowed_pages':[0],'slides':[{'page':0,'texts':{str(box.shape_id):'new'}}]})
        self.assertEqual(len(prs.slides),1);self.assertEqual(before,{str(x.shape_id):geometry(x) for x in prs.slides[0].shapes});self.assertEqual(prs.slides[0].shapes[0].text,'new');self.assertTrue(validate(out)['passed'])
    def test_sparse_placeholder_idx(self):
        p=Presentation(self.template);p.slide_layouts[0].placeholders[1]._element.ph.set('idx','12');p.save(self.template)
        prs,out=self.run_plan({'slides':[{'layout':0,'title':'Title','body':['Text']}]});self.assertEqual(prs.slides[0].placeholders[12].text,'Text');self.assertTrue(validate(out)['passed'])
    def test_creator_produces_usable_layouts(self):
        assets=self.root/'assets';(assets/'logos').mkdir(parents=True);(assets/'palettes').mkdir()
        Image.new('RGB',(100,100),'purple').save(assets/'logos/logo.png');Image.new('RGB',(100,100),'green').save(assets/'palettes/colors.png')
        template=self.root/'brand.pptx';mapping=create(assets,template,'Arial')
        self.assertEqual(len(mapping),7);self.assertEqual(len(Presentation(template).slides),0)
        output=self.root/'brand-deck.pptx';generate(template,{'mode':'layouts','slides':[{'layout':2,'title':'Brand','body':['Content']}]},output);self.assertTrue(validate(output)['passed'])
    def test_all_styles_generate(self):
        for name in ['minimal','tech-blue','academic','festive-red','traditional','casual']:
            with self.subTest(style=name):
                out=self.root/(name+'.pptx');create_deck(name,{'slides':[{'kind':'content','title':'Test','body':['Readable']} ]},out);self.assertTrue(validate(out)['passed'])
    def test_overflow_warning(self):
        prs,out=self.run_plan({'slides':[{'layout':0,'title':'X'*3000}]});self.assertTrue(any(x['code']=='text_fit_estimate' for x in validate(out)['issues']))
    def test_sample_small_font_does_not_split_short_line(self):
        from pptx.util import Inches,Pt
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
        box=s.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(.5))
        box.text='原文';box.text_frame.paragraphs[0].runs[0].font.size=Pt(12)
        p.save(self.template)
        value='使用模板原有的小字号填充完整文字内容'
        prs,out=self.run_plan({'mode':'sample-pages','allowed_pages':[0],'slides':[{'page':0,'texts':{str(box.shape_id):value}}]})
        self.assertEqual(len(prs.slides),1);self.assertEqual(prs.slides[0].shapes[0].text,value)
    def test_sample_preserves_master_visibility(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s._element.set('showMasterSp','0');p.save(self.template)
        prs,out=self.run_plan({'mode':'sample-pages','allowed_pages':[0],'slides':[{'page':0,'texts':{}}]})
        self.assertEqual(prs.slides[0]._element.get('showMasterSp'),'0')

if __name__=='__main__':unittest.main()
