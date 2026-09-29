"""Exercise picture replacement, source mapping, and optional PowerPoint notes."""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from common import geometry
from template_engine import generate
from write_source_report import write_report


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "source.pptx"
        self.output = self.root / "output.pptx"
        self.replacement = self.root / "replacement.png"
        original_image = self.root / "original.png"
        Image.new("RGB", (400, 200), "red").save(original_image)
        Image.new("RGB", (400, 200), "blue").save(self.replacement)
        prs = Presentation()
        for i in range(2):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(original_image), Inches(1), Inches(1), Inches(4), Inches(2))
        prs.save(self.source)
        self.picture_id = str(Presentation(self.source).slides[0].shapes[0].shape_id)
        self.plan = {
            "edits": [{"page": 0, "images": {self.picture_id: str(self.replacement)}, "notes": "讲解蓝色插图。"}],
            "speaker_notes": {"1": "第二页讲稿。"},
            "sources": [{"page": 0, "title": "第一页来源", "url": "https://example.org/one", "accessed": "2026-09-29"},
                        {"page": 1, "title": "第二页来源", "url": "https://example.org/two", "accessed": "2026-09-29"}],
            "image_sources": [{"page": 0, "shape_id": self.picture_id, "title": "测试插图", "origin": "AI-generated"}],
        }

    def test_picture_replaced_in_place_and_sources_stay_on_their_pages(self):
        before = Presentation(self.source).slides[0].shapes[0]
        audit = generate(self.source, self.plan, self.output)
        after = Presentation(self.output).slides[0].shapes[0]
        self.assertEqual(geometry(before), geometry(after))
        self.assertNotEqual(before.image.blob, after.image.blob)
        self.assertEqual(after.image.blob, self.replacement.read_bytes())
        self.assertEqual(audit["slides"][0]["replaced_images"], [self.picture_id])
        report = write_report(self.output, self.plan, audit).read_text(encoding="utf-8")
        first, second = report.split("### 第 2 页")
        self.assertIn("第一页来源", first)
        self.assertNotIn("第二页来源", first)
        self.assertIn("第二页来源", second)
        self.assertNotIn("第一页来源", second)

    @unittest.skipUnless(os.name == "nt" and os.getenv("SMART_PPT_TEST_POWERPOINT") == "1", "PowerPoint COM opt-in")
    def test_powerpoint_finalizer_writes_every_slide_note(self):
        generate(self.source, self.plan, self.output)
        plan_path = self.root / "plan.json"
        plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                        str(SCRIPTS / "finalize_powerpoint.ps1"), "-Pptx", str(self.output),
                        "-Plan", str(plan_path), "-Audit", str(self.output) + ".audit.json"],
                       timeout=90, capture_output=True, text=True, encoding="utf-8", errors="replace")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        prs = Presentation(self.output)
        self.assertIn("讲解蓝色插图", prs.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn("第二页讲稿", prs.slides[1].notes_slide.notes_text_frame.text)
        self.assertTrue((self.root / "信息源报告.md").is_file())


if __name__ == "__main__":
    unittest.main()
