"""Inspect layouts and authored sample pages without modifying the source."""
import argparse
from pathlib import Path
from pptx import Presentation
from common import geometry,walk,write_json

def inspect(path):
    prs=Presentation(str(path));layouts=[]
    for mi,master in enumerate(prs.slide_masters):
        for i,l in enumerate(master.slide_layouts):
            ph=[{'idx':s.placeholder_format.idx,'type':s.placeholder_format.type.name.lower(),
                 'name':s.name,'geometry_emu':geometry(s)} for s in l.placeholders]
            layouts.append({'master_index':mi,'index':i,'name':l.name,'placeholders':ph})
    pages=[]
    for i,s in enumerate(prs.slides):
        pages.append({'index':i,'name':s.name or f'Slide {i+1}',
            'shape_count':len(list(walk(s.shapes))),
            'text_shapes':[{'shape_id':x.shape_id,'name':x.name,'placeholder':x.is_placeholder,
                           'text':x.text[:160],'geometry_emu':geometry(x)}
                          for x in walk(s.shapes) if x.has_text_frame]})
    return {'template':Path(path).name,'slide_size':{'width_inches':prs.slide_width/914400,
            'height_inches':prs.slide_height/914400},'layouts':layouts,'sample_pages':pages}
def table(data):
    lines=['| Master | Index | Layout | Placeholder idx:type |','|---|---|---|---|']
    for l in data['layouts']:
        ph=', '.join(f"{p['idx']}:{p['type']}" for p in l['placeholders']) or '(no editable placeholder)'
        lines.append(f"| {l['master_index']} | {l['index']} | {l['name'].replace('|','/')} | {ph} |")
    return '\n'.join(lines)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('template');p.add_argument('--output');p.add_argument('--table',action='store_true')
    a=p.parse_args();d=inspect(a.template)
    if a.output:write_json(a.output,d)
    if a.table:print(table(d))
    else:
        import json;print(json.dumps(d,ensure_ascii=False,indent=2))
