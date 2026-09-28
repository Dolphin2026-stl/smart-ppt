"""Default to complete-deck edits; explicit layout/sample extraction remains available."""
import argparse
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from common import *
from layout_inspector import inspect,table

def remove_slides(prs,keep=()):
    # Isolated private API: python-pptx has no public delete-slide API.
    for i,sid in reversed(list(enumerate(prs.slides._sldIdLst))):
        if i not in keep:
            prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)

def clone_page(prs,source):
    dest=prs.slides.add_slide(source.slide_layout)
    for key,value in source._element.attrib.items():dest._element.set(key,value)
    for s in list(dest.shapes):s._element.getparent().remove(s._element)
    relmap={}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(('/slideLayout','/notesSlide')):continue
        relmap[rel.rId]=dest.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,rel.is_external)
    for s in source.shapes:
        elem=deepcopy(s._element)
        for node in elem.iter():
            for key,val in list(node.attrib.items()):
                if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and val in relmap:node.set(key,relmap[val])
        dest.shapes._spTree.insert_element_before(elem,'p:extLst')
    if source._element.cSld.bg is not None:
        dest._element.cSld.insert(0,deepcopy(source._element.cSld.bg))
    return dest

def usable(layout):
    return [p for p in layout.placeholders if p.placeholder_format.type not in IGNORED_TYPES]
def choose_layout(prs,item,allowed):
    candidates=[]
    for mi,idx in allowed:
        l=prs.slide_masters[mi].slide_layouts[idx];ph=usable(l)
        titles=[s for s in ph if s.placeholder_format.type in TITLE_TYPES]
        pics=[s for s in ph if s.placeholder_format.type==PP_PLACEHOLDER.PICTURE]
        bodies=[s for s in ph if s.has_text_frame and s not in titles and s not in pics]
        if item.get('title') and not titles:continue
        if item.get('body') and not bodies:continue
        if item.get('images') and not pics:continue
        candidates.append((abs(len(bodies)-len(item.get('body',[])))+abs(len(pics)-len(item.get('images',[]))),mi,idx))
    if not candidates:raise ValueError('No compatible selected layout. Inspect sample_pages for non-placeholder templates.')
    _,mi,idx=min(candidates);return mi,idx

def generate(template,plan,output):
    if plan.get('mode','full-deck')=='full-deck':
        from full_deck_engine import generate_full
        return generate_full(template,plan,output)
    ensure_new_output(template,output);prs=Presentation(str(template));audit={'mode':plan.get('mode','layouts'),'geometry_changes':[],'slides':[]}
    items=plan.get('slides',[])
    if not items:raise ValueError('slides must not be empty')
    if audit['mode']=='sample-pages':
        allowed=plan.get('allowed_pages')
        if not allowed:raise ValueError('sample-pages requires explicit allowed_pages (zero based)')
        originals=list(prs.slides)
        if any(not isinstance(i,int) or i<0 or i>=len(originals) for i in allowed):raise ValueError('Invalid page index')
        for item in items:
            idx=item['page']
            if idx not in allowed:raise ValueError('Page not in allowed_pages')
            source=originals[idx]
            # Fail closed on features whose cross-part relationships need deeper copying.
            for rel in source.part.rels.values():
                if rel.reltype.endswith(('/slide','/chart','/oleObject','/package')):raise ValueError('Sample page has linked slides/charts/OLE; use PowerPoint-native copying after review')
            mapping=item.get('texts',{})
            shapes={str(s.shape_id):s for s in walk(source.shapes)}
            for key,val in mapping.items():
                if key not in shapes or not shapes[key].has_text_frame:raise ValueError(f'Unknown text shape_id {key}')
            pieces={k:chunks(v,min(int(item.get('max_chars',400)),max(text_capacity(shapes[k]),len(shapes[k].text)))) for k,v in mapping.items()}
            count=max([len(v) for v in pieces.values()] or [1])
            for n in range(count):
                s=clone_page(prs,source);before={str(x.shape_id):geometry(x) for x in walk(s.shapes)}
                for x in walk(s.shapes):
                    key=str(x.shape_id)
                    if key in pieces:set_text(x,pieces[key][n] if n<len(pieces[key]) else '')
                if before!={str(x.shape_id):geometry(x) for x in walk(s.shapes)}:raise RuntimeError('Geometry changed')
                audit['slides'].append({'source_page':idx,'baseline':before,'filled_shapes':list(mapping)})
        remove_slides(prs,range(len(originals),len(prs.slides)))
    elif audit['mode']=='layouts':
        allowed=[tuple(x) for x in plan.get('allowed_layouts',[])]
        if not allowed:
            allowed=[(mi,i) for mi,m in enumerate(prs.slide_masters) for i,_ in enumerate(m.slide_layouts)]
        for mi,i in allowed:
            if mi<0 or i<0:raise ValueError('Indices must be nonnegative')
            prs.slide_masters[mi].slide_layouts[i]
        remove_slides(prs)
        for item in items:
            explicit=item.get('layout')
            if explicit is not None:
                mi,idx=(explicit if isinstance(explicit,list) else [0,explicit])
                if (mi,idx) not in allowed:raise ValueError('Layout not in allowed_layouts')
            else:mi,idx=choose_layout(prs,item,allowed)
            layout=prs.slide_masters[mi].slide_layouts[idx]
            ph=usable(layout)
            titles=[x for x in ph if x.placeholder_format.type in TITLE_TYPES]
            pics=[x for x in ph if x.placeholder_format.type==PP_PLACEHOLDER.PICTURE]
            bodies=[x for x in ph if x.has_text_frame and x not in titles and x not in pics]
            raw=item.get('body',[])
            if isinstance(raw,str):raw=[raw]
            images=list(item.get('images',[]))
            if item.get('title') and not titles:raise ValueError('Layout has no title placeholder')
            if raw and not bodies:raise ValueError('Layout has no body placeholder')
            if images and not pics:raise ValueError('Layout has no picture placeholder; refusing overlay')
            limit=min([text_capacity(x) for x in bodies] or [400])
            blocks=[p for b in raw for p in chunks(b,min(int(item.get('max_chars',400)),limit))]
            count=max(1,math.ceil(len(blocks)/max(1,len(bodies))),math.ceil(len(images)/max(1,len(pics))))
            for page in range(count):
                s=prs.slides.add_slide(layout);before={str(x.placeholder_format.idx):geometry(x) for x in s.placeholders}
                if titles:set_text(s.placeholders[titles[0].placeholder_format.idx],item.get('title',''),titles[0])
                for p,text in zip(bodies,blocks[page*len(bodies):(page+1)*len(bodies)]):set_text(s.placeholders[p.placeholder_format.idx],text,p)
                for p,img in zip(pics,images[page*len(pics):(page+1)*len(pics)]):s.placeholders[p.placeholder_format.idx].insert_picture(str(img))
                after={str(x.placeholder_format.idx):geometry(x) for x in s.placeholders}
                if before!=after:raise RuntimeError('Placeholder geometry changed')
                audit['slides'].append({'master_index':mi,'layout':idx,'reason':'explicit selection' if explicit is not None else 'matched required title/body/picture slots; overflow paginated','baseline_placeholders':before})
    else:raise ValueError('mode must be layouts or sample-pages')
    save(prs,output);write_json(str(output)+'.audit.json',audit);return audit
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('template');p.add_argument('plan');p.add_argument('output')
    a=p.parse_args();print(table(inspect(a.template)));generate(a.template,read_json(a.plan),a.output)
