"""Build editable themed layouts and populate them through the template engine."""
import argparse
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from common import *
from template_engine import generate,remove_slides
from layout_inspector import inspect
ROOT=Path(__file__).resolve().parents[1]

def load_style(name):
    path=Path(name)
    if not path.is_file():path=ROOT/'assets'/'styles'/(name+'.json')
    d=read_json(path)
    for k in ('primary','secondary','accent','background','text'):
        val=d['colors'][k]
        if len(val)!=6 or any(c not in '0123456789abcdefABCDEF' for c in val):raise ValueError('Colors must be six hex digits')
    for k in ('cover','agenda','content','two-column','three-column','chart','ending'):
        if k not in d['layouts']:raise ValueError('Missing layout '+k)
    return d

def rect(slide,box,color,kind=MSO_SHAPE.RECTANGLE):
    s=slide.shapes.add_shape(kind,*[Inches(v) for v in box]);s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color);s.line.fill.background();s._element.spPr.append(OxmlElement('a:effectLst'))
    for node in s._element.xpath('./p:style/a:effectRef'):node.set('idx','0')
    return s

def picture_fit(slide,path,box):
    from PIL import Image
    with Image.open(path) as im:w,h=im.size
    x,y,bw,bh=box;scale=min(bw/w,bh/h)
    return slide.shapes.add_picture(str(path),Inches(x+(bw-w*scale)/2),Inches(y+(bh-h*scale)/2),width=Inches(w*scale),height=Inches(h*scale))

def build_template(style,output,brand=None):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    brand=brand or {};events=[];colors=style['colors'];layouts={}
    for i,(name,spec) in enumerate(style['layouts'].items()):
        if i>=len(prs.slide_layouts):raise ValueError('At most 11 layouts supported')
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(colors['background'])
        if brand.get('background') and name in ('cover','ending'):
            picture_fit(slide,brand['background'],[8.9,.7,3.8,5.7])
        deco=style['decorations']['kind']
        rect(slide,[.55,.55,.08,6.35],colors['accent'])
        if deco=='grid':
            for x in [9.9,10.65,11.4,12.15]:rect(slide,[x,.5,.012,6.5],colors['secondary'])
            for y in [1,2,3,4,5,6]:rect(slide,[9.5,y,3.25,.012],colors['secondary'])
        elif deco=='seal':rect(slide,[11.85,5.95,.65,.65],colors['accent'])
        elif deco=='round':rect(slide,[11.5,5.7,1,1],colors['accent'],MSO_SHAPE.OVAL)
        elif deco=='band':rect(slide,[.63,6.9,12.1,.12],colors['primary'])
        if brand.get('logo'):picture_fit(slide,brand['logo'],[11.6,.25,1.05,.6])
        for idx,slot in enumerate(spec):
            role=slot['role'];shape=slide.shapes.add_textbox(*[Inches(v) for v in slot['box']]);shape.name=f'{role}-{idx}'
            shape.text_frame.word_wrap=True
            shape.text_frame.margin_left=Inches(.03);shape.text_frame.margin_right=Inches(.03)
            shape.text_frame.margin_top=Inches(.02);shape.text_frame.margin_bottom=Inches(.02)
            p=shape.text_frame.paragraphs[0];p.font.name=style['fonts']['title' if role=='title' else 'body'];p.font.size=Pt(slot.get('size',style['fonts']['sizes']['title' if role=='title' else 'body']))
            p.font.color.rgb=RGBColor.from_string(colors['text']);p.font.bold=role=='title'
            ppr=p._p.get_or_add_pPr();ppr.set('algn','l');ppr.insert(0,OxmlElement('a:buNone'))
            for tag in ('a:ea','a:cs'):
                node=OxmlElement(tag);node.set('typeface',style['fonts']['title' if role=='title' else 'body']);p.font._rPr.append(node)
            # Promote authored blank-slide boxes to real layout placeholders.
            nv=shape._element.nvSpPr;nv.cNvSpPr.attrib.pop('txBox',None)
            ph=OxmlElement('p:ph');ph.set('type',{'title':'title','body':'body','picture':'pic'}[role]);ph.set('idx',str(idx));nv.nvPr.append(ph)
        layout=prs.slide_layouts[i];layout.name=name
        for old in list(layout.shapes):old._element.getparent().remove(old._element)
        for shape in slide.shapes:
            elem=deepcopy(shape._element)
            # Images need layout-owned relationships.
            for node in elem.iter():
                for key,val in list(node.attrib.items()):
                    if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                        rel=slide.part.rels[val];node.set(key,layout.part.relate_to(rel.target_part,rel.reltype))
            layout.shapes._spTree.insert_element_before(elem,'p:extLst')
            events.append({'layout':i,'shape':shape.name,'geometry_emu':geometry(shape),'reason':'initial authoring; no source template modified'})
        bg=layout._element.cSld.bg
        if bg is not None:layout._element.cSld.remove(bg)
        layout._element.cSld.insert(0,deepcopy(slide._element.cSld.bg));layouts[name]=[0,i]
    remove_slides(prs);save(prs,output)
    write_json(str(output)+'.layouts.json',inspect(output));write_json(str(output)+'.creation.json',{'geometry_changes':events,'style':style,'brand':brand})
    return layouts

def create_deck(style_name,content,output):
    style=load_style(style_name);template=Path(output).with_suffix('.template.pptx');mapping=build_template(style,template)
    slides=[]
    for item in content['slides']:
        row=dict(item);kind=row.pop('kind','content');row['layout']=mapping[kind];slides.append(row)
    return generate(template,{'mode':'layouts','allowed_layouts':list(mapping.values()),'slides':slides},output)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('content');p.add_argument('output');p.add_argument('--style',default='minimal');a=p.parse_args();create_deck(a.style,read_json(a.content),a.output)
