"""Edit a complete PPTX package, preserving untouched parts byte-for-byte."""
import hashlib,posixpath,shutil,zipfile
from pathlib import Path
from copy import deepcopy
from lxml import etree as E
from pptx import Presentation
from common import ensure_new_output,geometry,walk,set_text,write_json
R='http://schemas.openxmlformats.org/package/2006/relationships'
O='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
C='http://schemas.openxmlformats.org/package/2006/content-types'
def relpath(part):return posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')
def resolve(part,target):return posixpath.normpath(posixpath.join(posixpath.dirname(part),target)).lstrip('/')
def xml(data):return E.fromstring(data)
def dump(node):return E.tostring(node,encoding='UTF-8',xml_declaration=True,standalone=True)
def generate_full(template,plan,output):
    ensure_new_output(template,output)
    if 'slides' in plan:raise ValueError('Full-deck mode uses edits, delete_pages or sequence. Use explicit mode=layouts for new decks.')
    prs=Presentation(str(template));original=list(prs.slides);n=len(original)
    def index(i):
        if not isinstance(i,int) or isinstance(i,bool) or not 0<=i<n:raise ValueError(f'Invalid source page {i}')
        return i
    deleted=plan.get('delete_pages',[])
    for i in deleted:index(i)
    if 'sequence' in plan and deleted:raise ValueError('Use sequence or delete_pages, not both')
    sequence=plan.get('sequence',[i for i in range(n) if i not in deleted])
    if not sequence:raise ValueError('Output must retain at least one page')
    for item in sequence:index(item['source_page'] if isinstance(item,dict) else item)
    edits={}
    for item in plan.get('edits',[]):
        i=index(item['page'])
        if i in edits:raise ValueError('Combine edits for the same page')
        edits[i]=item.get('texts',{})
    chosen=[x['source_page'] if isinstance(x,dict) else x for x in sequence]
    if set(edits)-set(chosen):raise ValueError('Cannot edit a deleted page')
    baseline=[{str(s.shape_id):geometry(s) for s in walk(slide.shapes)} for slide in original]
    audit={'mode':'full-deck','source_sha256':hashlib.sha256(Path(template).read_bytes()).hexdigest(),'source_slide_count':n,'geometry_changes':[],'deleted_pages':sorted(set(range(n))-set(chosen)),'slides':[]}
    if not edits and 'sequence' not in plan and not deleted:
        Path(output).parent.mkdir(parents=True,exist_ok=True);shutil.copy2(template,output)
        audit['slides']=[{'source_page':i,'baseline':baseline[i],'filled_shapes':[]} for i in range(n)]
        audit['unchanged_copy']=True;write_json(str(output)+'.audit.json',audit);return audit
    with zipfile.ZipFile(template) as z:data={p:z.read(p) for p in z.namelist()}
    source_data=dict(data)
    types=xml(data['[Content_Types].xml']);type_by={x.get('PartName').lstrip('/'):x for x in types if x.tag.endswith('Override')}
    pres=xml(data['ppt/presentation.xml']);pr=xml(data['ppt/_rels/presentation.xml.rels']);ids=pres.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst')
    original_ids=list(ids);rel_by={x.get('Id'):x for x in pr};parts=[str(s.part.partname).lstrip('/') for s in original]
    counter=[0]
    def clone(part,mapping):
        if part in mapping:return mapping[part]
        while True:
            counter[0]+=1;dest=posixpath.join(posixpath.dirname(part),'smartppt_'+str(counter[0])+'_'+posixpath.basename(part))
            if dest not in data:break
        mapping[part]=dest;data[dest]=source_data[part]
        if part in type_by:
            item=deepcopy(type_by[part]);item.set('PartName','/'+dest);types.append(item)
        rp=relpath(part)
        if rp in data:
            rels=xml(source_data[rp])
            for rel in rels:
                if rel.get('TargetMode')=='External':continue
                target=resolve(part,rel.get('Target'));kind=rel.get('Type').rsplit('/',1)[-1]
                if target in mapping:new=mapping[target]
                elif kind in ('slideLayout','slideMaster','theme','image','audio','video','media','slide'):new=target
                else:new=clone(target,mapping)
                rel.set('Target',posixpath.relpath(new,posixpath.dirname(dest)))
            data[relpath(dest)]=dump(rels)
        return dest
    for child in list(ids):ids.remove(child)
    seen=set();next_id=max(int(x.get('id')) for x in original_ids)+1;used_rids=set(rel_by)
    for entry in sequence:
        i=entry['source_page'] if isinstance(entry,dict) else entry
        texts=dict(edits.get(i,{}));texts.update(entry.get('texts',{}) if isinstance(entry,dict) else {})
        source=original[i];part=parts[i]
        if i in seen:
            part=clone(part,{});sid=deepcopy(original_ids[i]);sid.set('id',str(next_id));next_id+=1
            rid='rIdSmartPpt'+str(next_id)
            while rid in used_rids:rid+='x'
            used_rids.add(rid);rel=E.SubElement(pr,'{'+R+'}Relationship',Id=rid,Type=O+'/slide',Target=posixpath.relpath(part,'ppt'));sid.set('{'+O+'}id',rid)
        else:sid=deepcopy(original_ids[i])
        seen.add(i);ids.append(sid)
        if texts:
            # Fresh source proxy prevents one duplicate's text edit leaking into another.
            from pptx.slide import Slide
            working=Slide(deepcopy(source._element),source.part)
            shape_map={str(s.shape_id):s for s in walk(working.shapes)}
            for key,value in texts.items():
                if key not in shape_map or not shape_map[key].has_text_frame:raise ValueError(f'No editable text shape {key} on page {i}')
                set_text(shape_map[key],value)
            if baseline[i]!={str(s.shape_id):geometry(s) for s in walk(working.shapes)}:raise RuntimeError('Geometry changed')
            data[part]=dump(working._element)
        audit['slides'].append({'source_page':i,'baseline':baseline[i],'filled_shapes':list(texts),'duplicated':chosen.count(i)>1})
    active={x.get('{'+O+'}id') for x in ids}
    for rel in list(pr):
        if rel.get('Type')==O+'/slide' and rel.get('Id') not in active:pr.remove(rel)
    removed={parts[i] for i in range(n) if i not in chosen}
    for sid in ids:
        part=resolve('ppt/presentation.xml',next(x for x in pr if x.get('Id')==sid.get('{'+O+'}id')).get('Target'))
        if relpath(part) in data:
            for rel in xml(data[relpath(part)]):
                if rel.get('TargetMode')!='External' and rel.get('Type')==O+'/slide' and resolve(part,rel.get('Target')) in removed:raise ValueError('A retained page links to a deleted page; resolve the link before deletion')
    if chosen!=list(range(n)):
        data['ppt/presentation.xml']=dump(pres);data['ppt/_rels/presentation.xml.rels']=dump(pr)
        if 'docProps/app.xml' in data:
            app=xml(data['docProps/app.xml'])
            for node in app:
                if node.tag.endswith('}Slides'):node.text=str(len(sequence))
            data['docProps/app.xml']=dump(app)
    if removed:
        reachable=set();queue=['']
        while queue:
            part=queue.pop();rp='_rels/.rels' if not part else relpath(part)
            if part:reachable.add(part)
            if rp in data:
                reachable.add(rp)
                for rel in xml(data[rp]):
                    if rel.get('TargetMode')=='External':continue
                    target=resolve(part,rel.get('Target'))
                    if target not in reachable:reachable.add(target);queue.append(target)
        data={k:v for k,v in data.items() if k in reachable or k=='[Content_Types].xml'}
        for item in list(types):
            if item.tag.endswith('Override') and item.get('PartName').lstrip('/') not in data:types.remove(item)
    if counter[0] or removed:data['[Content_Types].xml']=dump(types)
    Path(output).parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        for name,blob in data.items():z.writestr(name,blob)
    audit['output_slide_count']=len(sequence);write_json(str(output)+'.audit.json',audit);return audit
