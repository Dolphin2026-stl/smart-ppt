"""Structural checks plus conservative text fit estimates; not a rendering substitute."""
import argparse,io,math,unicodedata
from pathlib import Path
from PIL import Image
from pptx import Presentation
from common import *

def validate(path,audit_path=None):
    prs=Presentation(str(path));issues=[]
    def issue(level,code,slide,shape,message):issues.append(dict(level=level,code=code,slide=slide,shape=shape,message=message))
    ap=Path(audit_path or str(path)+'.audit.json');audit=read_json(ap) if ap.exists() else None
    if audit and len(audit.get('slides',[]))!=len(prs.slides):issue('error','audit_slide_count',0,0,'Audit does not match slide count')
    for n,slide in enumerate(prs.slides,1):
        baseline=audit['slides'][n-1] if audit and n<=len(audit.get('slides',[])) else {}
        for s in walk(slide.shapes):
            if 'baseline' in baseline:
                expected=baseline['baseline'].get(str(s.shape_id))
                if expected is None or geometry(s)!=expected:issue('error','baseline_geometry',n,s.shape_id,'Sample-page shape differs from source baseline')
            if s.is_placeholder:
                idx=s.placeholder_format.idx
                l=slide.slide_layout.placeholders.get(idx)
                if l is None:issue('warning','missing_layout_placeholder',n,s.shape_id,'No matching layout placeholder idx')
                else:
                    if geometry(s)!=geometry(l):
                        level='warning' if 'baseline' in baseline else 'error'
                        issue(level,'layout_geometry',n,s.shape_id,'Geometry differs from layout; may be intentional in source sample page')
                    if s._element.xpath('./p:spPr/a:xfrm'):
                        issue('info','explicit_transform',n,s.shape_id,'Explicit xfrm exists; equal geometry is allowed, inspect audit for intent')
                expected=baseline.get('baseline_placeholders',{}).get(str(idx))
                if expected is not None and geometry(s)!=expected:issue('error','baseline_geometry',n,s.shape_id,'Placeholder geometry changed after filling')
            if s.has_text_frame and s.text.strip():
                tf=s.text_frame;w=max(1,(s.width-tf.margin_left-tf.margin_right)/12700);h=max(1,(s.height-tf.margin_top-tf.margin_bottom)/12700)
                need=0
                for p in tf.paragraphs:
                    size=max([x.font.size.pt for x in p.runs if x.font.size] or [p.font.size.pt if p.font.size else 24])
                    units=sum(1 if unicodedata.east_asian_width(c) in 'WF' else .55 for c in p.text)
                    need+=max(1,math.ceil(units*size/w))*size*1.25
                if need>h*1.08:issue('warning','text_fit_estimate',n,s.shape_id,f'Estimated text height {need:.1f}pt > box {h:.1f}pt; render to confirm')
            if hasattr(s,'image'):
                try:
                    iw,ih=s.image.size;dx,dy=s.image.dpi
                    visible_w=iw*(1-s.crop_left-s.crop_right)/dx;visible_h=ih*(1-s.crop_top-s.crop_bottom)/dy
                    if visible_w<=0 or visible_h<=0:raise ValueError('invalid crop')
                    expected=visible_w/visible_h;actual=s.width/s.height
                    if abs(actual/expected-1)>.025:issue('warning','picture_aspect',n,s.shape_id,'Displayed aspect differs from cropped intrinsic aspect; possible stretching')
                except (ValueError,ZeroDivisionError,AttributeError) as e:issue('warning','picture_unchecked',n,s.shape_id,str(e))
        expected_ids=baseline.get('baseline',{})
        if expected_ids and set(expected_ids)!={str(s.shape_id) for s in walk(slide.shapes)}:issue('error','shape_set',n,0,'Sample-page shapes added or removed')
        expected_ph=baseline.get('baseline_placeholders',{})
        if expected_ph and set(expected_ph)!={str(s.placeholder_format.idx) for s in slide.placeholders}:issue('error','placeholder_set',n,0,'Placeholders added or removed')
    return {'file':Path(path).name,'slide_count':len(prs.slides),'passed':not any(x['level']=='error' for x in issues),'render_review_required':True,'limitations':['Text overflow is a heuristic; actual font metrics and rendering may differ.','Without a baseline audit, inherited-layout changes cannot be attributed.','Cropped pictures are compared after cropping; crop is not distortion.'],'issues':issues}
def markdown(report):
    lines=[f"# Validation: {report['file']}",f"Structural pass: {report['passed']}; slides: {report['slide_count']}",'','Visual review is still required.','','| Level | Slide | Shape | Check | Message |','|---|---|---|---|---|']
    lines += [f"| {x['level']} | {x['slide']} | {x['shape']} | {x['code']} | {x['message']} |" for x in report['issues']]
    return '\n'.join(lines)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('pptx');p.add_argument('--audit');p.add_argument('--report');p.add_argument('--strict',action='store_true');a=p.parse_args()
    d=validate(a.pptx,a.audit)
    if a.report:
        write_json(a.report,d);Path(a.report).with_suffix('.md').write_text(markdown(d),encoding='utf-8')
    print(markdown(d));raise SystemExit(1 if not d['passed'] or a.strict and any(x['level']=='warning' for x in d['issues']) else 0)
