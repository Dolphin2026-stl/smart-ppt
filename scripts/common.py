"""Shared deterministic IO and geometry utilities (inches in JSON, EMU in audit)."""
import json, math, unicodedata
from pathlib import Path
from copy import deepcopy
from pptx.enum.shapes import PP_PLACEHOLDER

TITLE_TYPES={PP_PLACEHOLDER.TITLE,PP_PLACEHOLDER.CENTER_TITLE,PP_PLACEHOLDER.VERTICAL_TITLE}
IGNORED_TYPES={PP_PLACEHOLDER.DATE,PP_PLACEHOLDER.FOOTER,PP_PLACEHOLDER.HEADER,PP_PLACEHOLDER.SLIDE_NUMBER}
def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))
def write_json(path, data):
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
def geometry(s):
    return [int(getattr(s,k) or 0) for k in ('left','top','width','height')]
def walk(shapes):
    for s in shapes:
        yield s
        if hasattr(s,'shapes'): yield from walk(s.shapes)
def set_text(shape,text,style_source=None):
    """Preserve first paragraph and run styling; never create or resize shapes."""
    if not isinstance(text,str): raise ValueError('Text must be a string')
    tf=shape.text_frame
    first=(style_source.text_frame if style_source is not None else tf).paragraphs[0]
    ppr=deepcopy(first._p.pPr) if first._p.pPr is not None else None
    rpr=deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    tf.clear()
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if p._p.pPr is not None: p._p.remove(p._p.pPr)
        if ppr is not None: p._p.insert(0,deepcopy(ppr))
        run=p.add_run();run.text=line
        if rpr is not None: run._r.insert(0,deepcopy(rpr))
def chunks(text,limit):
    if not isinstance(text,str): raise ValueError('Text must be a string')
    if limit<1: raise ValueError('max_chars must be positive')
    return [text[i:i+limit] for i in range(0,len(text),limit)] or ['']
def text_capacity(s,points=None):
    if not s.has_text_frame:return 0
    tf=s.text_frame
    if points is None:
        sizes=[run.font.size.pt for p in tf.paragraphs for run in p.runs if run.font.size]
        sizes += [p.font.size.pt for p in tf.paragraphs if p.font.size]
        points=max(sizes) if sizes else 24
    w=max(1,(s.width-tf.margin_left-tf.margin_right)/12700)
    h=max(1,(s.height-tf.margin_top-tf.margin_bottom)/12700)
    return max(1,int(w/points)*int(h/(points*1.3)))
def ensure_new_output(source,output):
    if Path(source).resolve()==Path(output).resolve():raise ValueError('Refusing to overwrite source template')
def save(prs,path):
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(str(path))
